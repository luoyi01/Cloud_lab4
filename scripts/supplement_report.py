"""Add mapped operation screenshots to an existing Word report, preserving edits.

Use a pre-supplement report as --source. This deliberately refuses a second
insertion, so an edited report is never silently regenerated or duplicated.
"""
import argparse
from datetime import datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

from build_word_report import add_text, normalize_package, set_font

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / 'screenshots/08-live-validation'
LOGS = ROOT / 'results/screenshot-validation'

GROUPS = [
 ('4.5.1 环境与容器化', [
  ('01-env', 'V01', '环境和工具实际版本', '核对 Ubuntu、内核、Docker 客户端与服务端、Compose、Kind 和 kubectl 的版本。本机使用现有 WSL2 与 Docker Desktop。'),
  ('02-build', 'V02', '独立构建并运行 Dockerfile 镜像', 'API 和 Web 分别由 Dockerfile 构建，输出真实镜像 ID；独立验证容器读取 Redis 并提供 HTTP 服务后清理。补拍标签 screenshot 与主部署标签 1.0 分开使用。'),
  ('03-compose-config', 'V03', '端口 环境变量 命名卷和内部发现', '8090 映射到 Web 的 80 端口；容器读到标题、版本和 Redis 主机配置。Redis 的 /data 使用命名卷，内部 DNS 能解析 redis，业务 GET 能读取共享存储。'),
  ('04-compose-before', 'V04 V09', 'Compose 原始部署与共享业务验证', '原始部署包含 3 个 API、Web 和 Redis。业务检查通过固定入口识别 3 个不同实例，写入标记并由各实例读回同一记录，访问计数递增。该图也是 Compose 删除对照的部署前状态。'),
 ]),
 ('4.5.2 Kubernetes 基础资源与配置', [
  ('05-k8s-resources', 'V05', '三节点集群和应用资源就绪', '3 个节点 Ready；API Deployment 为 3/3，Web 和 Redis 就绪。API/Web 的 Service 与 NodePort 提供服务发现和入口，Redis 的 StatefulSet 与 PVC 显示存储身份和 Bound 状态。'),
  ('06-k8s-config', 'V06', 'ConfigMap 和 Secret 引用及运行时注入', 'ConfigMap 中的标题、版本和 Redis 主机与容器 printenv 一致。Deployment 引用 cloudlab-secret/REDIS_PASSWORD；Secret 仅显示字段名和字节数。业务读回成功同时验证实际 Redis 认证可用。'),
 ]),
 ('4.5.3 Pod 自愈前中后组图', [
  ('07-heal-before', 'V07', '原始 3 副本和连续业务请求', '删除前确认 3 个 Ready Pod 与 Service 就绪端点一致。启动 45 秒连续业务 GET，每 0.2 秒请求一次，不做客户端重试；请求 CSV 单独保存。'),
  ('08-heal-inject', 'V07', '手动删除和替代 Pod 创建过程', '命令手动删除目标 Pod。输出同时包含旧 Pod 的 Terminating 状态和新 Pod 的创建状态，说明 Deployment 根据期望副本数触发替换。旧 UID 保留用于后续对照。'),
  ('09-heal-after', 'V07', '恢复到 3 副本和请求连续性', '恢复后 3 个 Ready Pod 与 Service 端点匹配，新 UID 与被删除 UID 不同。图中请求统计覆盖部署、删除和恢复全过程；错误数按实际 CSV 计算。原实验 3.711 秒的计时结论仍依据第五节原始记录。'),
 ]),
 ('4.5.4 两条路径扩缩容与资源组图', [
  ('10-scale-one', 'V08', '伸缩起点 1 副本与资源用量', 'Compose 与 Kubernetes 都调整到 1 个 API 副本，等待就绪后展示容器资源、Pod 数量和 Pod CPU/内存。'),
  ('11-scale-three', 'V08', '扩容到 3 副本与资源用量', '同一应用在两条路径中扩到 3 个 API 副本。docker stats 与 kubectl top 对照副本数量和实际用量；limit 为容量上限，不等于实际占用。'),
  ('12-scale-five', 'V08', '扩容到 5 副本与资源用量', '两条路径继续扩到 5 个 API，截图显示全部实例和指标。API 无本地业务状态，共享 Redis 保持数据一致。'),
  ('13-scale-back', 'V08', '缩回 3 副本与资源用量', '通过显式伸缩命令缩回 3 个 API，确认额外实例退出。组图展示 1→3→5→3 的实际操作；这里是补拍时的瞬时指标，第五节表格仍使用原固定负载下的五次观测中位数。'),
 ]),
 ('4.5.5 数据持久化前中后组图', [
  ('14-persist-before', 'V03 V10 B01', '重启前写入标记和确认存储', '分别向 Compose 与 Kubernetes 写入各自的新标记，保留记录 ID、内容和创建时间。命名卷与 PVC 是本组重启后读回的存储基础。'),
  ('15-persist-inject', 'V10 B01', 'Redis 容器重启与 StatefulSet Pod 重建', 'Compose 重启 Redis 容器，Kubernetes 删除 redis-0，输出记录删除命令和即时 Pod 状态；同名替代 Pod 就绪见恢复图。操作采用正常终止，不是突然断电测试。'),
  ('16-persist-after', 'V03 V10 B01', '恢复后读回原记录并确认 PVC', '先确认 redis-0 的 UID 已改变、新 Pod Ready 且未处于删除状态，再从两个公开业务入口读回原标记，记录完整相等，输出 identical_readback=True。PVC 保持 Bound，证明正常重启后的持久化；单实例 Redis 仍不具备主从故障转移。'),
 ]),
 ('4.5.6 Compose 删除对照与恢复', [
  ('17-compose-inject', 'V09', '删除容器后 15 秒仍为 2 副本', '部署前状态见 Compose 原始部署图。删除一个 API 容器，15 秒后仍只有 2 个，入口继续返回业务响应。restart 策略不会重建已删除容器，与 Pod 自愈组形成实际对照。'),
  ('18-compose-after', 'V09', '显式执行 Compose 恢复 3 副本', '执行 compose up --scale api=3 后恢复期望数量，展示最终 3 个 API 和覆盖删除、恢复阶段的连续请求统计。服务发现由 Compose 内部 DNS 和 Nginx 代理完成。'),
 ]),
 ('4.5.7 PDB 自愿驱逐前中后组图', [
  ('19-pdb-before', 'B01', '3 副本时的可驱逐预算', 'API 3 个副本健康，PDB minAvailable=2 允许预算为 1。该约束保护自愿驱逐路径。'),
  ('20-pdb-inject', 'B01', '2 副本时拒绝 Eviction', '将副本数改为 2，预算变为 0，再通过 policy/v1 Eviction 请求驱逐。API Server 拒绝请求，已有 Pod 仍存在；直接 delete 不受 PDB 阻止。'),
  ('21-pdb-after', 'B01', '恢复 3 副本和预算', '将副本恢复为 3，等待全部就绪，预算恢复为 1。完整组图同时体现配置、约束触发和最终恢复。'),
 ]),
 ('4.5.8 原始对比数据复核', [
  ('22-comparison', 'V09', '两条路径各五次热启动原始记录', '现场执行读取原始 JSON 的命令，显示每轮路径、轮次、耗时和原始 UTC 时间。这是数据复核截图，热启动测量发生在图中标明的原始日期；不把补拍时间当作测量时间。副本、恢复、配置和发现的机制比较见表 9。'),
 ]),
 ('4.5.9 API worker 故障前中后组图', [
  ('23-node-before', 'B02', '故障前节点与 Pod 分布', '先确认被停 worker 确实承载 Ready API，记录受影响 UID，启动 140 秒连续业务请求。Web、Redis 与本地 PV 留在存活 worker，本组仅验证 API 节点故障。'),
  ('24-node-inject', 'B02', '停止 worker 并观察 NotReady', '停止 cloudlab-worker2 节点容器，等待 Kubernetes 检测后显示 NotReady 与 Pod 状态。控制器发现失联需要时间，不能将这一窗口写成服务完全不中断。'),
  ('25-node-after', 'B02', '存活节点重新调度和故障节点恢复', '等待 3 个 Ready API 全部调度到存活 worker 后，重新启动故障 worker，确认三个节点 Ready。图中保留本轮请求成功和错误数；第五节原节点实验的 44 次错误仍保留。单机 Kind 不证明宿主机或 Redis 高可用。'),
 ]),
]


def paragraph(doc, text, style=None, size=11, bold=False, keep=False):
    p = doc.add_paragraph(style=style)
    add_text(p, text, size, bold)
    p.paragraph_format.keep_with_next = keep
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    return p


def matrix(doc):
    rows = [
        ('V01', '环境和工具实际版本', '图 4'),
        ('V02', 'Dockerfile 独立构建并正常运行', '图 5'),
        ('V03', '端口 环境变量 卷与通信生效', '图 6、17–19'),
        ('V04', 'Compose 多组件业务与共享数据', '图 2、7'),
        ('V05', 'Kind 集群 Deployment 和 Service', '图 3、8'),
        ('V06', 'ConfigMap 和 Secret 配置注入', '图 9'),
        ('V07', '手动删除 Pod 自动重建 连续请求', '图 10–12；原结果图 30'),
        ('V08', '伸缩前后数量与 CPU 内存用量', '图 13–16；原统计图 29'),
        ('V09', '同应用部署 副本 恢复 配置 发现 伸缩对比', '图 6–16、20–21、25；表 9'),
        ('V10', '正常重启后数据持久化', '图 17–19'),
        ('B01', '加分 StatefulSet PVC 和 PDB', '图 8、17–19、22–24'),
        ('B02', '加分 多节点 API worker 故障', '图 26–28；原结果图 31'),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [1.4, 8.0, 6.6]
    for col,w in zip(table.columns,widths): col.width = Cm(w)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        b=OxmlElement('w:'+edge)
        for key,val in [('val','single'),('sz','4'),('color','D9D9D9')]: b.set(qn('w:'+key),val)
        borders.append(b)
    table._tbl.tblPr.append(borders)
    repeat=OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,values in enumerate([('验证项','文档要求与验证内容','报告图号')] + rows):
        row=table.rows[0] if i==0 else table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell,value,w in zip(row.cells,values,widths):
            cell.width=Cm(w)
            p=cell.paragraphs[0]
            p.paragraph_format.space_before=Pt(4)
            p.paragraph_format.space_after=Pt(4)
            p.paragraph_format.line_spacing=1.15
            add_text(p,value,10.5,i==0)
            if i==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E9EEF2');cell._tc.get_or_add_tcPr().append(shd)
    return table


def build(source, output, anonymous=False):
    doc=Document(source)
    if any(p.text.startswith('4.5 操作截图') for p in doc.paragraphs):
        raise ValueError('Source already supplemented; select the pre-supplement Word backup.')
    if anonymous:
        identity=doc.paragraphs[1]
        identity.text='学号：________________    姓名：________________'
        for r in identity.runs: set_font(r.font,12)
    # Keep captions in reading order; existing analytical figures move after 25 new figures.
    for p in doc.paragraphs:
        replacement=re.sub(r'图\s*([456])(?!\d)', lambda m:'图 '+str(int(m[1])+25),p.text)
        replacement=re.sub(r'表\s*([5-9])(?!\d)', lambda m:'表 '+str(int(m[1])+1),replacement)
        if replacement.startswith('图 2 '): replacement += '（V04）'
        elif replacement.startswith('图 3 '): replacement += '（V05）'
        elif replacement.startswith('图 30 '): replacement += '（V07 原始结果复核）'
        elif replacement.startswith('图 31 '): replacement += '（B02 原始结果复核）'
        if replacement!=p.text:
            size=p.runs[0].font.size if p.runs else None
            p.text=replacement
            for r in p.runs: set_font(r.font,size.pt if size else None)
    for index in (1,2,4,5):
        shape=doc.inline_shapes[index]
        ratio=shape.height/shape.width
        width=min(16.0,17.5/ratio)
        shape.width=Cm(width)
        shape.height=Cm(width*ratio)
    inside_reflection=False
    for p in doc.paragraphs:
        if p.text.startswith('七、'): inside_reflection=True
        if p.text.startswith('八、'): inside_reflection=False
        if inside_reflection and p.style.name=='Normal':
            p.paragraph_format.line_spacing=1.35
            p.paragraph_format.space_after=Pt(5)
    anchor=next(p._p for p in doc.paragraphs if p.text.startswith('五、消融'))
    added=[]
    before=set(doc.element.body)
    paragraph(doc,'4.5 操作截图与验证项映射','Heading 2',12,True,True)
    paragraph(doc,'图 4–28 为补充操作截图，直接捕获本机实时命令执行界面。固定命令在 WSL2 中现场执行，页面显示命令、stdout/stderr 和 UTC 时间；界面由 scripts/live_validation.py 提供。各组按验证项组织，操作时间分别标注，不作为一次连续实验。补拍记录单独保存至 results/screenshot-validation/，原测量数据保留不变。')
    paragraph(doc,'截图按验证项编号；Pod 自愈、持久化、Compose 删除、PDB 和节点故障均给出操作前、注入过程、恢复后的对照。扩缩容按 1→3→5→3 展示。图 25 等需要读取原统计材料的步骤在图注中明确为结果复核；图 29–31 为原统计图与记录核对视图。')
    paragraph(doc,'表 5  操作截图覆盖索引',size=11,bold=True,keep=True)
    matrix(doc)
    figure=4
    manifest=[]
    for heading,records in GROUPS:
        paragraph(doc,heading,'Heading 2',12,True,True)
        for stage,validation,title,description in records:
            shot=SHOTS/f'{stage}.png'
            log=LOGS/f'{stage}.json'
            record=json.loads(log.read_text(encoding='utf-8'))
            assert record['status'].startswith('执行完成') and shot.is_file(),stage
            with Image.open(shot) as img: w,h=img.size
            width=min(16.0,17.5*w/h)
            p=doc.add_paragraph()
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_after=Pt(4)
            p.add_run().add_picture(str(shot),width=Cm(width))
            caption=f'图 {figure}  {validation}  {title}'
            p=paragraph(doc,caption,size=10.5,bold=True,keep=True)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            local=datetime.fromisoformat(record['utc']).astimezone(timezone(timedelta(hours=8)))
            paragraph(doc,description+f' 操作时间：{local:%Y-%m-%d %H:%M:%S}（北京时间）；命令和完整输出：results/screenshot-validation/{stage}.txt。',size=10.5)
            manifest.append({'figure':figure,'validation':validation,'stage':stage,'caption':caption,'description':description,
                             'utc':record['utc'],'screenshot':shot.relative_to(ROOT).as_posix(),
                             'log':log.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(shot.read_bytes()).hexdigest()})
            figure+=1
    added=[e for e in list(doc.element.body) if e not in before and e.tag!=qn('w:sectPr')]
    for e in added: anchor.addprevious(e)
    doc.core_properties.author=''
    doc.core_properties.last_modified_by=''
    doc.core_properties.comments=''
    for style in doc.styles:
        if hasattr(style,'font'): set_font(style.font)
    doc.save(output)
    normalize_package(output)
    (LOGS/'figure-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return manifest


def markdown(docx_path, output):
    doc=Document(docx_path)
    images={hashlib.sha256(p.read_bytes()).hexdigest():p for folder in (ROOT/'screenshots',ROOT/'docs/figures') for p in folder.rglob('*.png')}
    lines=[]
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    for item in doc.element.body:
        if item.tag==qn('w:p'):
            p=Paragraph(item,doc)
            text=p.text
            if p.style.name=='Title': text='# '+text
            elif p.style.name=='Heading 1': text='## '+text
            elif p.style.name=='Heading 2': text='### '+text
            if text: lines.extend([text,''])
            for drawing in item.findall('.//'+qn('a:blip')):
                part=doc.part.related_parts[drawing.get(qn('r:embed'))]
                source=images[hashlib.sha256(part.blob).hexdigest()]
                import os
                relative=Path(os.path.relpath(source,output.parent)).as_posix()
                lines.extend([f'![操作截图或分析图]({relative})',''])
        elif item.tag==qn('w:tbl'):
            t=Table(item,doc)
            rows=[[c.text.replace('\n','；').replace('|','\\|') for c in r.cells] for r in t.rows]
            lines.append('| '+' | '.join(rows[0])+' |')
            lines.append('| '+' | '.join(['---']*len(rows[0]))+' |')
            lines.extend('| '+' | '.join(r)+' |' for r in rows[1:])
            lines.append('')
    output.write_text('\n'.join(lines).rstrip()+'\n',encoding='utf-8')


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--anonymous',action='store_true')
    parser.add_argument('--markdown',type=Path)
    args=parser.parse_args()
    result=build(args.source,args.output,args.anonymous)
    if args.markdown: markdown(args.output,args.markdown)
    print(f'Added {len(result)} mapped real operation screenshots; all editable fonts SimSun.')
